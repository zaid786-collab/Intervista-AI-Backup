"""
Resume Document Parser Service for Intervista AI
Supports PDF and DOCX document extraction with in-memory processing.
Validates file formats, sizes, detects password protection, and detects scanned/image-only PDFs.
"""

import io
import re
from typing import Dict, Any, Tuple


MAX_FILE_SIZE_BYTES = 10 * 1024 * 1024  # 10 MB
MIN_FILE_SIZE_BYTES = 50  # 50 bytes


def clean_extracted_text(raw_text: str) -> str:
    """Cleans up raw extracted text from PDFs and DOCX files."""
    if not raw_text:
        return ""
    # Normalize excessive null or control chars
    cleaned = re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f\x7f]", " ", raw_text)
    # Normalize multiple newlines into double newlines
    cleaned = re.sub(r"\n{3,}", "\n\n", cleaned)
    # Normalize repeated spaces and tabs
    cleaned = re.sub(r"[ \t]{2,}", " ", cleaned)
    return cleaned.strip()


def parse_pdf_document(raw_bytes: bytes, filename: str) -> Dict[str, Any]:
    """
    Extracts text and metadata from a PDF file using pypdf with pdfplumber fallback.
    Detects password protection and scanned/image-only documents.
    """
    import pypdf

    stream = io.BytesIO(raw_bytes)
    try:
        reader = pypdf.PdfReader(stream)
    except Exception as e:
        return {
            "success": False,
            "error": f"Could not read PDF structure: {str(e)}. The document may be corrupted.",
            "is_corrupted": True,
        }

    # 1. Check for password encryption
    if reader.is_encrypted:
        return {
            "success": False,
            "error": "This PDF is password-protected. Please unlock or export an unencrypted PDF before uploading.",
            "is_encrypted": True,
        }

    page_count = len(reader.pages)
    if page_count == 0:
        return {
            "success": False,
            "error": "The uploaded PDF document contains 0 pages.",
            "page_count": 0,
        }

    extracted_pages = []
    for idx, page in enumerate(reader.pages):
        try:
            page_text = page.extract_text() or ""
            if page_text.strip():
                extracted_pages.append(page_text.strip())
        except Exception:
            pass

    full_text = "\n\n".join(extracted_pages).strip()

    # 2. If pypdf extracted very little text, try pdfplumber fallback
    if len(full_text.split()) < 20:
        try:
            import pdfplumber
            stream.seek(0)
            plumber_pages = []
            with pdfplumber.open(stream) as pdf:
                for p in pdf.pages:
                    txt = p.extract_text() or ""
                    if txt.strip():
                        plumber_pages.append(txt.strip())
            if plumber_pages:
                alt_text = "\n\n".join(plumber_pages).strip()
                if len(alt_text.split()) > len(full_text.split()):
                    full_text = alt_text
        except Exception:
            pass

    full_text = clean_extracted_text(full_text)
    word_count = len(full_text.split())

    # 3. Detect scanned / image-based PDFs
    if word_count < 20:
        return {
            "success": False,
            "is_scanned": True,
            "error": (
                "Scanned or image-only PDF detected. The document contains pages but no selectable text. "
                "Applicant Tracking Systems (ATS) require readable text layers to parse resumes. "
                "Please upload a searchable text-based PDF or run OCR before uploading."
            ),
            "page_count": page_count,
            "word_count": word_count,
        }

    return {
        "success": True,
        "text": full_text,
        "file_name": filename,
        "file_type": "pdf",
        "page_count": page_count,
        "word_count": word_count,
        "char_count": len(full_text),
    }


def parse_docx_document(raw_bytes: bytes, filename: str) -> Dict[str, Any]:
    """
    Extracts text and metadata from a Word (.docx) document using python-docx.
    Traverses paragraphs and tables to extract comprehensive content.
    """
    import docx

    stream = io.BytesIO(raw_bytes)
    try:
        doc = docx.Document(stream)
    except Exception as e:
        return {
            "success": False,
            "error": f"Could not open Word document: {str(e)}. Please ensure it is a valid .docx file.",
            "is_corrupted": True,
        }

    paragraphs = []
    # Extract from main paragraphs
    for p in doc.paragraphs:
        txt = p.text.strip()
        if txt:
            paragraphs.append(txt)

    # Extract from tables (many ATS resumes use tabular structures for skills or dates)
    for table in doc.tables:
        for row in table.rows:
            row_cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if row_cells:
                # Deduplicate repeated merged cell text in the same row
                seen = set()
                unique_cells = []
                for c in row_cells:
                    if c not in seen:
                        seen.add(c)
                        unique_cells.append(c)
                paragraphs.append(" | ".join(unique_cells))

    full_text = clean_extracted_text("\n\n".join(paragraphs))
    word_count = len(full_text.split())

    if word_count < 20:
        return {
            "success": False,
            "error": (
                "The Word document appears to be empty or contains fewer than 20 readable words. "
                "Please upload a complete resume document."
            ),
            "word_count": word_count,
        }

    return {
        "success": True,
        "text": full_text,
        "file_name": filename,
        "file_type": "docx",
        "page_count": max(1, word_count // 350),
        "word_count": word_count,
        "char_count": len(full_text),
    }


def parse_resume_file(raw_bytes: bytes, filename: str) -> Dict[str, Any]:
    """
    Universal entry point to validate and parse resume files (PDF or DOCX).
    Enforces file size and type boundaries.
    """
    # 1. Size Validation
    if not raw_bytes or len(raw_bytes) < MIN_FILE_SIZE_BYTES:
        return {
            "success": False,
            "error": "The uploaded file is empty or contains no readable data.",
        }

    if len(raw_bytes) > MAX_FILE_SIZE_BYTES:
        return {
            "success": False,
            "error": f"File size ({len(raw_bytes) // (1024 * 1024)} MB) exceeds the maximum allowed size of 10 MB.",
        }

    clean_filename = (filename or "resume.pdf").strip().lower()

    # 2. Format Validation & Delegation
    if clean_filename.endswith(".pdf"):
        return parse_pdf_document(raw_bytes, filename)
    elif clean_filename.endswith(".docx"):
        return parse_docx_document(raw_bytes, filename)
    elif clean_filename.endswith(".doc"):
        return {
            "success": False,
            "error": (
                "Legacy .doc format is not supported directly due to proprietary binary encoding. "
                "Please save/export your resume as a modern .docx or .pdf file and re-upload."
            ),
        }
    else:
        return {
            "success": False,
            "error": "Unsupported file format. Please upload a PDF (.pdf) or Word document (.docx).",
        }
